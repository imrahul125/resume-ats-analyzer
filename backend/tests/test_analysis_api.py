"""Direct route tests use a fake session so they never touch a real database."""

from io import BytesIO
import unittest
from uuid import uuid4

try:
    from fastapi import HTTPException
    from fastapi.testclient import TestClient
    from starlette.datastructures import UploadFile
    from app.api.routes import _validate_upload, analyze_resume
    from app.core.database import get_db
    from app.main import app
except ModuleNotFoundError as exc:  # Workspace utility Python may not include app dependencies.
    raise unittest.SkipTest(f"API test dependencies are unavailable: {exc.name}") from exc

from docx import Document


class FakeSession:
    def __init__(self) -> None:
        self.records = []
        self.pending = []
        self.committed = False

    def add_all(self, records) -> None:
        self.pending.extend(records)
        self.records.extend(records)

    def add(self, record) -> None:
        self.pending.append(record)
        self.records.append(record)

    def flush(self) -> None:
        for record in self.pending:
            if hasattr(record, "id") and record.id is None:
                record.id = uuid4()
        self.pending.clear()

    def commit(self) -> None:
        self.committed = True

    def rollback(self) -> None:
        self.committed = False


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Summary", level=1)
    document.add_paragraph("Python engineer with six years building backend systems and APIs.")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Built REST APIs using Python and FastAPI for production services.")
    document.add_heading("Education", level=1)
    document.add_paragraph("Bachelor of Science")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python, FastAPI, PostgreSQL")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


class AnalysisRouteTests(unittest.TestCase):
    def test_upload_validation_checks_type_and_content(self) -> None:
        with self.assertRaises(HTTPException) as unsupported:
            _validate_upload("resume.exe", "application/octet-stream", b"data")
        self.assertEqual(unsupported.exception.status_code, 415)

        with self.assertRaises(HTTPException) as fake_pdf:
            _validate_upload("resume.pdf", "application/pdf", b"not a PDF")
        self.assertEqual(fake_pdf.exception.status_code, 415)

    def test_analyze_saves_results_without_saving_resume_text(self) -> None:
        upload = UploadFile(
            filename="resume.docx",
            file=BytesIO(make_docx()),
            headers={"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        )
        session = FakeSession()
        result = analyze_resume(
            resume_file=upload,
            job_description=(
                "Required qualifications:\nPython and FastAPI are required.\n"
                "Preferred:\nDocker is nice to have."
            ),
            job_title="Backend Engineer",
            company="Example",
            db=session,
        )

        self.assertTrue(session.committed)
        self.assertGreaterEqual(result.overall_score, 0)
        self.assertLessEqual(result.overall_score, 100)
        self.assertEqual(len(result.metrics), 7)
        self.assertEqual({item.name: item.status for item in result.skills}["Python"], "matched")
        self.assertTrue(any(item.category == "missing_requirement" for item in result.recommendations))
        self.assertTrue(any(record.__class__.__name__ == "Analysis" for record in session.records))
        resume_record = next(record for record in session.records if record.__class__.__name__ == "Resume")
        self.assertFalse(hasattr(resume_record, "raw_text"))
        self.assertFalse(hasattr(resume_record, "content"))

    def test_multipart_http_request_returns_serialized_analysis(self) -> None:
        session = FakeSession()
        app.dependency_overrides[get_db] = lambda: session
        self.addCleanup(app.dependency_overrides.clear)
        response = TestClient(app).post(
            "/api/v1/analyze",
            data={"job_description": "Required: Python is required. Preferred: Docker is a plus."},
            files={
                "resume_file": (
                    "resume.docx",
                    make_docx(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        self.assertEqual(response.status_code, 201, response.text)
        body = response.json()
        self.assertGreaterEqual(body["overall_score"], 0)
        self.assertLessEqual(body["overall_score"], 100)
        self.assertEqual(len(body["metrics"]), 7)
        self.assertEqual(body["skills"][0]["name"], "Python")
        self.assertTrue(session.committed)


if __name__ == "__main__":
    unittest.main()
