"""Fast unit coverage for parsing helpers and transparent scoring rules."""

from io import BytesIO
import unittest

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from app.services.jd_analyzer import (
    classify_skill_in_resume,
    extract_job_requirements,
    normalize_skill,
)
from app.services.resume_parser import detect_sections, normalize_text, parse_docx, parse_pdf
from app.services.scoring import WEIGHTS, analyze_text


class ResumeParserTests(unittest.TestCase):
    def test_normalize_text_keeps_paragraphs_and_removes_control_characters(self) -> None:
        self.assertEqual(normalize_text("  Hello\x00   world\n\nNext  "), "Hello world\nNext")

    def test_parse_docx_extracts_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_heading("Experience", level=1)
        document.add_paragraph("Built APIs with Python")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Education"
        table.cell(0, 1).text = "Bachelor of Science"
        stream = BytesIO()
        document.save(stream)

        extracted = parse_docx(stream.getvalue())
        self.assertIn("Built APIs with Python", extracted)
        self.assertIn("Education | Bachelor of Science", extracted)
        self.assertEqual(detect_sections(extracted), ["experience"])

    def test_blank_pdf_has_no_extractable_text(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        stream = BytesIO()
        writer.write(stream)
        self.assertEqual(parse_pdf(stream.getvalue()), "")

    def test_parse_pdf_extracts_selectable_text(self) -> None:
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 40 700 Td (Python Engineer Resume) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        stream = BytesIO()
        writer.write(stream)

        self.assertIn("Python Engineer Resume", parse_pdf(stream.getvalue()))


class SkillAndScoringTests(unittest.TestCase):
    def test_skill_aliases_normalize_to_canonical_names(self) -> None:
        self.assertEqual(normalize_skill("Postgres"), "PostgreSQL")
        self.assertEqual(normalize_skill("JS"), "JavaScript")
        self.assertIsNone(normalize_skill("made-up technology"))

    def test_required_and_preferred_skills_use_job_description_context(self) -> None:
        requirements = extract_job_requirements(
            "Required qualifications:\nPython and FastAPI are required.\n"
            "Preferred:\nDocker is nice to have."
        )
        self.assertEqual(
            {item.name: item.requirement for item in requirements},
            {"Python": "required", "FastAPI": "required", "Docker": "preferred"},
        )

    def test_multiword_requirement_can_be_partial(self) -> None:
        self.assertEqual(classify_skill_in_resume("Built web features with React", "React Native"), "partial")

    def test_score_is_weighted_from_visible_factors(self) -> None:
        resume = (
            "SUMMARY\nPython engineer with several years building production services.\n"
            "EXPERIENCE\n- Built RESTful APIs using Python and FastAPI.\n"
            "EDUCATION\nBachelor of Science\n"
            "SKILLS\nPython, FastAPI, PostgreSQL"
        )
        job = "Required qualifications:\nPython and FastAPI are required.\nPreferred:\nDocker is nice to have.\nBachelor's degree."
        result = analyze_text(resume, job)
        weighted_score = round(sum(item["score"] * item["weight"] for item in result["metrics"]))

        self.assertEqual(result["overall_score"], weighted_score)
        self.assertEqual(sum(weight for _, weight in WEIGHTS.values()), 1.0)
        self.assertEqual(len(result["metrics"]), 7)
        states = {item["name"]: item["status"] for item in result["skills"]}
        self.assertEqual(states["Python"], "matched")
        self.assertEqual(states["Docker"], "missing")
        self.assertTrue(any("Docker" in item["message"] for item in result["recommendations"]))


if __name__ == "__main__":
    unittest.main()
