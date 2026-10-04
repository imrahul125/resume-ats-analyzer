"""Extract readable text from supported resume formats without saving uploads."""

from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader


SECTION_ALIASES = {
    "summary": {"summary", "professional summary", "profile", "career profile", "objective"},
    "experience": {"experience", "work experience", "professional experience", "employment history"},
    "education": {"education", "academic background", "educational background"},
    "skills": {"skills", "technical skills", "core competencies", "areas of expertise"},
    "projects": {"projects", "personal projects", "selected projects"},
    "certifications": {"certifications", "certificates", "licenses and certifications"},
    "achievements": {"achievements", "accomplishments", "awards and honors", "honors"},
    "internships": {"internships", "internship experience"},
}
_HEADER_TO_SECTION = {
    re.sub(r"[^a-z0-9 ]", "", alias.lower()).strip(): section
    for section, aliases in SECTION_ALIASES.items()
    for alias in aliases
}


def normalize_text(text: str) -> str:
    """Normalize Unicode and whitespace while preserving paragraph boundaries."""
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", text)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def parse_pdf(content: bytes) -> str:
    """Return text extracted from a text-based PDF, or a safe validation error."""
    try:
        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ValueError("This PDF is password-protected. Upload an unlocked PDF.")
        if len(reader.pages) > 50:
            raise ValueError("The PDF has too many pages. Upload a resume with 50 pages or fewer.")
        extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("We couldn't read this PDF. Try exporting it again and re-uploading it.") from exc
    return normalize_text(extracted)


def parse_docx(content: bytes) -> str:
    """Extract paragraph and table text from a DOCX held only in memory."""
    try:
        document = Document(BytesIO(content))
        chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    chunks.append(" | ".join(cells))
    except (BadZipFile, PackageNotFoundError, KeyError, ValueError) as exc:
        raise ValueError("We couldn't read this DOCX. Open it in Word and save it as a new .docx file.") from exc
    except Exception as exc:
        raise ValueError("We couldn't read this DOCX. Try saving a fresh .docx copy and re-uploading it.") from exc
    return normalize_text("\n".join(chunks))


def parse_resume(filename: str, content: bytes) -> str:
    """Choose a parser from the allowed filename extension."""
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(content)
    if suffix == ".docx":
        return parse_docx(content)
    raise ValueError("Unsupported file type. Upload a PDF or DOCX resume.")


def extract_sections(text: str) -> dict[str, str]:
    """Split recognizable resume sections; text before the first heading is ignored."""
    sections: dict[str, list[str]] = defaultdict(list)
    current: str | None = None
    for line in text.splitlines():
        heading = re.sub(r"[^a-z0-9 ]", "", line.lower()).strip()
        section = _HEADER_TO_SECTION.get(heading)
        if section:
            current = section
        elif current:
            sections[current].append(line.strip())
    return {name: "\n".join(lines).strip() for name, lines in sections.items()}


def detect_sections(text: str) -> list[str]:
    """Return standard section names detected in the resume, in document order."""
    detected: list[str] = []
    for line in text.splitlines():
        heading = re.sub(r"[^a-z0-9 ]", "", line.lower()).strip()
        section = _HEADER_TO_SECTION.get(heading)
        if section and section not in detected:
            detected.append(section)
    return detected
